from __future__ import annotations

import asyncio
import hashlib
from io import BytesIO
from uuid import uuid4

import httpx
import pytest

from app.models import ItemStatus
from app.seed import DDS_LICENSE_ID, MODULE_ALPHA_ID
from app.services.file_guard_client import FileGuardClient, FileGuardUnavailableError
from app.config import Settings
from tests.test_api import make_client, auth
from tests.test_file_guard_client import upload


@pytest.fixture
def register(tmp_path):
    client = make_client(tmp_path)
    employee = auth(client, "employee", "employee")
    request = client.post("/requests", json={"unit_id": MODULE_ALPHA_ID}, headers=employee).json()
    repo = client.app.state.repo
    for index in range(121):
        repo.create("req_items", {"id": str(uuid4()), "request_id": request["id"], "dds_id": DDS_LICENSE_ID, "name": f"Line {index}", "sum_plan": index + 1, "sum_fact": 0, "status": ItemStatus.on_review, "analytics_1": "A" if index % 2 else "B"})
    return client, employee, request


def test_summary_and_paged_details_preserve_legacy_results(register):
    client, headers, request = register
    filters = {"request_id": request["id"]}
    legacy = client.get("/approval-register", params=filters, headers=headers).json()
    summary = client.post("/approval-register/query", json={"filters": filters}, headers=headers)
    assert summary.status_code == 200, summary.text
    data = summary.json()
    assert "summary_items" not in data
    assert data["aggregates"] == legacy["aggregates"]
    assert len(summary.content) < len(client.get("/approval-register", params=filters, headers=headers).content) / 4
    ids = []
    for page in (1, 2, 3):
        response = client.post("/approval-register/query", json={"mode": "rows", "filters": filters, "page": page}, headers=headers)
        assert response.status_code == 200, response.text
        rows = response.json()["items"]
        assert len(rows) <= 50
        ids.extend(row["id"] for row in rows)
        assert all(row["status_context"] is not None for row in rows)
    assert ids == [row["id"] for row in legacy["summary_items"]]
    assert len(set(ids)) == 121


def test_filters_sorting_and_last_page_apply_to_entire_scope(register):
    client, headers, request = register
    body = {"filters": {"request_id": request["id"]}, "columns": {"analytics_1": ["A"]}, "sort": {"column": "requested", "direction": "desc"}}
    summary = client.post("/approval-register/query", json=body, headers=headers).json()
    assert summary["aggregates"]["total_rows"] == 60
    assert len(summary["matched_item_ids"]) == 60
    rows = client.post("/approval-register/query", json={**body, "mode": "rows", "page": 999}, headers=headers).json()
    assert rows["pagination"]["page"] == 2
    assert len(rows["items"]) == 10
    assert [row["requested_sum"] for row in rows["items"]] == list(range(20, 0, -2))


def test_facets_and_group_loading_do_not_expose_hidden_requests(register):
    client, _, request = register
    economist = auth(client, "economist", "economist")
    for mode in ("summary", "rows", "facets", "groups", "selection"):
        response = client.post("/approval-register/query", json={"mode": mode, "filters": {"request_id": request["id"]}}, headers=economist)
        assert response.status_code == 200, response.text
        assert request["id"] not in str(response.json())


def test_guard_spools_large_processed_output_and_rewinds(monkeypatch):
    content = b"x" * (2 * 1024 * 1024)
    original_client = httpx.AsyncClient
    def handler(request):
        return httpx.Response(200, content=content, headers={"content-type": "application/pdf", "X-File-Guard-Action": "accepted", "X-File-Guard-Output-Name": "file.pdf", "X-File-Guard-Source-Sha256": "a" * 64, "X-File-Guard-Output-Sha256": hashlib.sha256(content).hexdigest()})
    monkeypatch.setattr(httpx, "AsyncClient", lambda **kwargs: original_client(transport=httpx.MockTransport(handler), **kwargs))
    file = upload()
    result = asyncio.run(FileGuardClient(Settings()).process(file))
    try:
        assert result.content == b""
        assert result.content_stream._rolled
        assert result.output_size_bytes == len(content)
        assert result.open_content().read() == content
        assert file.file.tell() == 0
    finally:
        result.close()
    assert result.content_stream.closed


def test_guard_rejects_oversized_processed_response(monkeypatch):
    original_client = httpx.AsyncClient
    monkeypatch.setattr(httpx, "AsyncClient", lambda **kwargs: original_client(transport=httpx.MockTransport(lambda request: httpx.Response(200, content=b"x" * 2048)), **kwargs))
    guard = FileGuardClient(Settings())
    guard.max_size_bytes = 1024
    with pytest.raises(FileGuardUnavailableError):
        asyncio.run(guard.process(upload()))


def test_cursor_history_has_no_duplicates_at_identical_timestamps(register):
    client, headers, request = register
    repo = client.app.state.repo
    user = client.get("/auth/me", headers=headers).json()
    for index in range(15):
        repo.create("req_logs", {"req_id": request["id"], "user_id": user["id"], "created_at": "2026-01-01T00:00:00+00:00", "log": {"action": "item_updated", "index": index}})
    expected = client.get(f'/requests/{request["id"]}/logs', headers=headers).json()
    rows, cursor = [], None
    while True:
        response = client.get(f'/requests/{request["id"]}/logs', params={"page_size": 5, **({"before": cursor} if cursor else {})}, headers=headers)
        assert response.status_code == 200, response.text
        page = response.json()
        rows += page["items"]
        cursor = page["next_cursor"]
        if cursor is None:
            break
    keys = [(row["source"], row["id"]) for row in rows]
    assert len(keys) == len(set(keys)) == len(expected)
    assert set(keys) == {(row["source"], row["id"]) for row in expected}


def test_items_and_notifications_keep_legacy_and_bounded_modes(register):
    client, headers, request = register
    response = client.get(f'/requests/{request["id"]}/items', params={"page_size": 50, "page": 3}, headers=headers)
    assert response.status_code == 200, response.text
    assert len(response.json()["items"]) == 21
    assert response.json()["pagination"]["total_items"] == 121
    assert isinstance(client.get('/notifications', headers=headers).json(), list)
    assert client.get('/notifications', params={"page_size": 50}, headers=headers).json() == {"items": [], "next_cursor": None}


def test_chat_pages_keep_latest_unread_and_reply_outside_page(register):
    client, headers, request = register
    repo = client.app.state.repo
    repo.update("requests", request["id"], {"status": "on_review"})
    chat = client.get(f'/requests/{request["id"]}/chat', headers=headers).json()
    stamp = "2026-01-01T00:00:00+00:00"
    ids = []
    for index in range(15):
        identity = str(uuid4())
        ids.append(identity)
        repo.create("chat_messages", {"id": identity, "chat_id": chat["id"], "sender_id": None, "is_system": True, "text": str(index), "created_at": stamp})
    ids.sort()
    repo.update("chat_messages", ids[-1], {"reply_to": ids[0]})
    first = client.get(f'/chats/{chat["id"]}', params={"page_size": 5}, headers=headers).json()
    assert len(first["messages"]) == 5
    assert first["unread_count"] == 15
    assert first["messages"][-1]["reply_preview"]["id"] == ids[0]
    rows, cursor = first["messages"], first["next_cursor"]
    while cursor:
        page = client.get(f'/chats/{chat["id"]}', params={"page_size": 5, "before": cursor}, headers=headers).json()
        assert page["last_message"]["id"] == ids[-1]
        assert page["unread_count"] == 15
        rows += page["messages"]
        cursor = page["next_cursor"]
    assert len(rows) == len({row["id"] for row in rows}) == 15


def test_export_uses_server_selection_without_url_ids(register):
    client, headers, request = register
    from openpyxl import load_workbook
    response = client.post('/approval-register/export', headers=headers, json={"query": {"filters": {"request_id": request["id"]}, "columns": {"analytics_1": ["A"]}}, "include_files": False})
    assert response.status_code == 200, response.text[:200] if response.status_code != 200 else ""
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames
    workbook.close()
